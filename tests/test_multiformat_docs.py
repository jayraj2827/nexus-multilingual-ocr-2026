import io
from pathlib import Path
import numpy as np
import openpyxl
from PIL import Image
import pytest
from pptx import Presentation
from docx import Document

from nexusocr.contracts.formats import FormatCategory, FormatResolver
from nexusocr.contracts.input import ProcessingOptions
from nexusocr.contracts.results import DocumentResult
from nexusocr.features.batch import BatchDocumentService
from nexusocr.features.document_ocr.forms import FormDataExtractor
from nexusocr.features.document_ocr.service import DocumentOCRService
from nexusocr.processors.office import DocxProcessor, PresentationProcessor, SpreadsheetProcessor
from nexusocr.processors.text_markup import TextMarkupProcessor


def test_docx_processor(tmp_path: Path):
    doc_path = tmp_path / 'test.docx'
    doc = Document()
    doc.add_heading('NexusOCR Architecture', level=1)
    doc.add_paragraph('This document outlines the multi-format capability.')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Format'
    table.cell(0, 1).text = 'Mode'
    table.cell(1, 0).text = 'DOCX'
    table.cell(1, 1).text = 'Native'
    doc.save(str(doc_path))

    service = DocumentOCRService()
    result = service.process_document(str(doc_path))
    assert isinstance(result, DocumentResult)
    assert result.total_pages == 1
    assert 'NexusOCR Architecture' in result.full_markdown
    assert 'Format' in result.full_markdown
    assert 'DOCX' in result.full_markdown


def test_spreadsheet_processor_xlsx(tmp_path: Path):
    xlsx_path = tmp_path / 'financial.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Summary'
    ws.append(['Quarter', 'Revenue', 'Profit'])
    ws.append(['Q1', ',000', ',500'])
    ws.append(['Q2', ',000', ',000'])
    wb.save(str(xlsx_path))

    service = DocumentOCRService()
    result = service.process_document(str(xlsx_path))
    assert isinstance(result, DocumentResult)
    assert result.total_pages == 1
    assert 'Quarter' in result.full_markdown
    assert ',000' in result.full_markdown
    assert len(result.pages[0].tables) == 1


def test_spreadsheet_processor_csv(tmp_path: Path):
    csv_path = tmp_path / 'data.csv'
    csv_path.write_text('Item,Qty,Price\nWidget,10,.00\nGadget,5,.00\n', encoding='utf-8')

    service = DocumentOCRService()
    result = service.process_document(str(csv_path))
    assert isinstance(result, DocumentResult)
    assert 'Widget' in result.full_markdown
    assert '.00' in result.full_markdown
    assert len(result.pages[0].tables) == 1


def test_presentation_processor_pptx(tmp_path: Path):
    pptx_path = tmp_path / 'deck.pptx'
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = 'NexusOCR Vision'
    slide.placeholders[1].text = 'Intelligent Document Processing'
    prs.save(str(pptx_path))

    service = DocumentOCRService()
    result = service.process_document(str(pptx_path))
    assert isinstance(result, DocumentResult)
    assert result.total_pages == 1
    assert 'NexusOCR Vision' in result.full_markdown


def test_text_and_html_processor(tmp_path: Path):
    # Plain text
    txt_path = tmp_path / 'memo.txt'
    txt_path.write_text('Important memorandum regarding OCR release.', encoding='utf-8')
    service = DocumentOCRService()
    res_txt = service.process_document(str(txt_path))
    assert 'Important memorandum' in res_txt.full_markdown

    # HTML
    html_path = tmp_path / 'table.html'
    html_path.write_text('<html><body><h1>Report</h1><table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table></body></html>', encoding='utf-8')
    res_html = service.process_document(str(html_path))
    assert 'Report' in res_html.full_markdown
    assert len(res_html.pages[0].tables) == 1


def test_multipage_tiff_ocr(tmp_path: Path):
    tiff_path = tmp_path / 'multi.tiff'
    img1 = Image.new('RGB', (200, 200), color=(255, 255, 255))
    img2 = Image.new('RGB', (200, 200), color=(240, 240, 240))
    img1.save(str(tiff_path), save_all=True, append_images=[img2])

    service = DocumentOCRService()
    result = service.process_document(str(tiff_path))
    assert isinstance(result, DocumentResult)
    assert result.total_pages == 2
    assert "Visual Summary & Caption" in result.full_markdown


def test_image_processing_and_captioning(tmp_path: Path):
    img_path = tmp_path / 'sample_doc.png'
    # Create image with high-contrast text-like blocks
    arr = np.ones((300, 400, 3), dtype=np.uint8) * 255
    arr[50:80, 50:350] = 0   # dark block resembling heading
    arr[120:150, 50:350] = 0 # dark block resembling text line
    img = Image.fromarray(arr)
    img.save(str(img_path))

    service = DocumentOCRService()
    result = service.process_document(str(img_path))
    assert isinstance(result, DocumentResult)
    assert result.total_pages == 1
    assert "Visual Summary & Caption" in result.full_markdown
    assert "Dimensions" in result.full_markdown
    assert len(result.pages[0].regions) > 0


def test_image_ocr_text_extraction(tmp_path: Path):
    from PIL import ImageDraw
    img_path = tmp_path / 'text_bill.png'
    img = Image.new('RGB', (500, 200), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((20, 30), "INVOICE #9988", fill=(0, 0, 0))
    draw.text((20, 80), "TOTAL DUE: $450.00", fill=(0, 0, 0))
    img.save(str(img_path))

    service = DocumentOCRService()
    result = service.process_document(str(img_path))
    assert isinstance(result, DocumentResult)
    assert "Visual Summary & Caption" in result.full_markdown
    assert "Transcribed Content" in result.full_markdown
    assert "INVOICE" in result.full_markdown
    assert len(result.pages[0].regions) >= 2


def test_pdf_embedded_image_extraction_and_ocr(tmp_path: Path):
    import fitz
    from PIL import ImageDraw
    pdf_path = tmp_path / 'doc_with_image.pdf'
    
    # 1. Create embedded image
    img = Image.new('RGB', (250, 120), color=(245, 245, 245))
    draw = ImageDraw.Draw(img)
    draw.text((15, 30), "DIAGRAM 1", fill=(0, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    img_bytes = buf.getvalue()

    # 2. Create PDF with text and insert image
    doc = fitz.open()
    page = doc.new_page(width=600, height=600)
    page.insert_text((50, 50), "Quarterly Financial Analysis Report")
    page.insert_text((50, 80), "The following diagram illustrates performance metrics:")
    
    rect = fitz.Rect(50, 120, 300, 240)
    page.insert_image(rect, stream=img_bytes)
    doc.save(str(pdf_path))
    doc.close()

    service = DocumentOCRService()
    result = service.process_document(str(pdf_path))
    assert isinstance(result, DocumentResult)
    assert "Quarterly Financial Analysis" in result.full_markdown
    assert "Embedded Image #1" in result.full_markdown
    assert "Visual Caption" in result.full_markdown
    assert "Executive Summary" not in result.full_markdown
    assert "Document Intelligence Analysis" not in result.full_markdown
    assert any(r.category == "figure" for r in result.pages[0].regions)


def test_docx_with_embedded_image(tmp_path: Path):
    from PIL import ImageDraw
    docx_path = tmp_path / 'memo_with_img.docx'

    img = Image.new('RGB', (300, 100), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((20, 30), "ATTACHED RECEIPT", fill=(0, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)

    doc = Document()
    doc.add_heading('Expense Report', level=1)
    doc.add_paragraph('Please see attached proof of payment below:')
    doc.add_picture(buf)
    doc.save(str(docx_path))

    service = DocumentOCRService()
    result = service.process_document(str(docx_path))
    assert isinstance(result, DocumentResult)
    assert "Expense Report" in result.full_markdown
    assert "Embedded Image #1" in result.full_markdown
    assert "Visual Caption" in result.full_markdown
    assert any(r.category == "figure" for r in result.pages[0].regions)


def test_pptx_with_embedded_image(tmp_path: Path):
    from PIL import ImageDraw
    pptx_path = tmp_path / 'presentation_with_img.pptx'

    img = Image.new('RGB', (300, 100), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((20, 30), "KEY RESULT AREA", fill=(0, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(buf, 0, 0)
    prs.save(str(pptx_path))

    service = DocumentOCRService()
    result = service.process_document(str(pptx_path))
    assert isinstance(result, DocumentResult)
    assert result.total_pages == 1
    assert "Slide Image #1" in result.full_markdown
    assert "Visual Caption" in result.full_markdown
    assert any(r.category == "figure" for r in result.pages[0].regions)


def test_api_image_upload_processing(tmp_path: Path):
    from fastapi.testclient import TestClient
    from app import app
    from PIL import ImageDraw

    img = Image.new('RGB', (400, 150), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((20, 40), "CERTIFICATE OF ACHIEVEMENT", fill=(0, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)

    client = TestClient(app)
    response = client.post(
        "/api/process",
        files={"file": ("award.png", buf, "image/png")}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["format"] == "IMAGE"
    assert "CERTIFICATE" in data["full_markdown"]




def test_form_data_extractor():
    text = '''
    INVOICE NUMBER: INV-2026-991
    DATE: 2026-09-13
    BILL TO: ACME Corporation
    TOTAL DUE: $12,450.00
    STATUS: Approved
    '''
    data = FormDataExtractor.extract_form_data(text)
    kv = data.get("key_value_pairs", {})
    assert len(kv) >= 3
    keys = {k.lower() for k in kv}
    assert any("invoice" in k for k in keys)

    currencies = data.get("entities", {}).get("currencies", [])
    assert any("12,450.00" in c for c in currencies)

    dates = data.get("entities", {}).get("dates", [])
    assert "2026-09-13" in dates


def test_batch_service_with_failure_isolation(tmp_path: Path):
    f1 = tmp_path / 'doc1.txt'
    f1.write_text('Batch item 1 content', encoding='utf-8')

    f2 = tmp_path / 'invalid_media.mp3'
    f2.write_bytes(b'bad media')

    f3 = tmp_path / 'doc2.txt'
    f3.write_text('Batch item 2 content', encoding='utf-8')

    service = DocumentOCRService()
    batch = BatchDocumentService(service)

    res = batch.process_batch([str(f1), str(f2), str(f3)])
    assert res.total_items == 3
    assert res.successful_count == 2
    assert res.failed_count == 1
    assert 'Audio file' in res.failures[0].error
