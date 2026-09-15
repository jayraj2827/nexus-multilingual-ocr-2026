"""
Test suite for PowerPoint presentation extraction and JSON serialization.
Verifies that embedded images do not leak raw bytes into structured_json
and that FastAPI jsonable_encoder succeeds with HTTP 200 payload structure.
"""

from pathlib import Path
import pytest
from fastapi.encoders import jsonable_encoder

from nexusocr.features.document_ocr.service import DocumentOCRService, _sanitize_for_json
from nexusocr.contracts.results import DocumentResult


def test_sanitize_for_json_removes_bytes():
    """Verifies that _sanitize_for_json replaces raw bytes and ignores embedded_images."""
    raw_dict = {
        "title": "Slide 1",
        "embedded_images": [b"\x89PNG\r\n\x1a\n\x00\x00"],
        "other_bytes": b"\x89PNGdata",
        "nested": {
            "blob": b"\x89binary",
            "text": "clean string",
        },
        "items": [1, "two", b"\xff\xd8rawjpeg"],
    }

    sanitized = _sanitize_for_json(raw_dict)

    # embedded_images key stripped from dicts
    assert "embedded_images" not in sanitized
    assert sanitized["title"] == "Slide 1"
    assert sanitized["other_bytes"] == "<8 bytes>"
    assert sanitized["nested"]["text"] == "clean string"
    assert sanitized["nested"]["blob"] == "<7 bytes>"
    assert sanitized["items"] == [1, "two", "<9 bytes>"]

    # jsonable_encoder must succeed without UnicodeDecodeError
    encoded = jsonable_encoder(sanitized)
    assert encoded["title"] == "Slide 1"


def test_pptx_file_processing_and_json_serialization():
    """Verifies processing of real PPTX with embedded images and ensures JSON encoding succeeds."""
    pptx_path = Path("data/uploads/Chapter_3_Data Link Layer.pptx")
    if not pptx_path.exists():
        pytest.skip(f"Test file {pptx_path} does not exist.")

    service = DocumentOCRService()
    result = service.process_document(str(pptx_path))

    assert isinstance(result, DocumentResult)
    assert result.format == "PRESENTATION"
    assert result.total_pages > 0
    assert len(result.pages) > 0

    # Ensure structured_json has slides with embedded_images_count, not raw bytes
    slides = result.structured_json.get("slides", [])
    assert len(slides) > 0
    for slide in slides:
        assert "embedded_images" not in slide
        assert "embedded_images_count" in slide
        assert isinstance(slide["embedded_images_count"], int)

    # Critical: jsonable_encoder must not raise UnicodeDecodeError: 'utf-8' codec can't decode byte 0x89
    encoded = jsonable_encoder(result)
    assert encoded["file_name"] == pptx_path.name
    assert "slides" in encoded["structured_json"]


def test_synthetic_pptx_with_png_image(tmp_path):
    """Verifies that a PPTX with embedded PNG (starting with byte 0x89) serializes cleanly."""
    import io
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Valid PNG image created via PIL
    pil_img = Image.new("RGB", (60, 60), color=(10, 100, 200))
    img_stream = io.BytesIO()
    pil_img.save(img_stream, format="PNG")
    img_stream.seek(0)

    # Verify that PNG starts with 0x89
    assert img_stream.getvalue()[0] == 0x89

    slide.shapes.add_picture(img_stream, Inches(1), Inches(1), width=Inches(2))

    pptx_file = tmp_path / "test_slide.pptx"
    prs.save(str(pptx_file))

    service = DocumentOCRService()
    result = service.process_document(str(pptx_file))

    assert result.format == "PRESENTATION"
    assert result.total_pages == 1
    assert "slides" in result.structured_json
    assert result.structured_json["slides"][0]["embedded_images_count"] == 1

    # Must serialize through jsonable_encoder without UnicodeDecodeError!
    encoded = jsonable_encoder(result)
    assert encoded["format"] == "PRESENTATION"
    assert encoded["structured_json"]["slides"][0]["embedded_images_count"] == 1

