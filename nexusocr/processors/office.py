"""
NexusOCR Office Document Media Processors.
Provides native structure, text, and table extraction for Word (.docx), Excel (.xlsx/.csv), and PowerPoint (.pptx).
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple


class DocxProcessor:
    """Extracts structured text, headings, tables, and embedded images from Word (.docx) files."""

    @staticmethod
    def extract_document(file_path: str) -> Dict[str, Any]:
        """
        Parses a .docx file and returns structured paragraphs, tables, and embedded image bytes.
        """
        import docx

        doc = docx.Document(file_path)
        markdown_sections: List[str] = []
        tables_data: List[Dict[str, Any]] = []
        embedded_images: List[bytes] = []

        # Extract paragraphs & headings
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name.lower() if p.style and p.style.name else ""
            if "heading 1" in style_name:
                markdown_sections.append(f"# {text}")
            elif "heading 2" in style_name:
                markdown_sections.append(f"## {text}")
            elif "heading 3" in style_name:
                markdown_sections.append(f"### {text}")
            elif "list" in style_name:
                markdown_sections.append(f"* {text}")
            else:
                markdown_sections.append(text)

        # Extract tables
        for idx, table in enumerate(doc.tables):
            grid: List[List[str]] = []
            for row in table.rows:
                row_cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
                grid.append(row_cells)

            if grid:
                header = grid[0]
                rows = grid[1:]
                header_line = "| " + " | ".join(header) + " |"
                sep_line = "| " + " | ".join(["---"] * max(1, len(header))) + " |"
                row_lines = ["| " + " | ".join(r) + " |" for r in rows]
                md_table = "\n".join([header_line, sep_line] + row_lines)

                tables_data.append({
                    "id": f"docx_table_{idx+1}",
                    "markdown": md_table,
                    "num_rows": len(grid),
                    "num_cols": len(header),
                    "grid": grid,
                })
                markdown_sections.append(md_table)

        # Extract embedded images from document relationships
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                try:
                    img_data = rel.target_part.blob
                    embedded_images.append(img_data)
                except Exception:
                    pass

        full_markdown = "\n\n".join(markdown_sections)
        return {
            "markdown": full_markdown,
            "tables": tables_data,
            "embedded_images": embedded_images,
            "total_paragraphs": len(doc.paragraphs),
            "total_tables": len(doc.tables),
        }


class SpreadsheetProcessor:
    """Extracts structured worksheets and tables from Excel (.xlsx) and CSV (.csv) files."""

    @staticmethod
    def extract_spreadsheet(file_path: str) -> Dict[str, Any]:
        """Parses spreadsheet files into structured sheets and Markdown tables."""
        ext = Path(file_path).suffix.lower()
        if ext == ".csv":
            return SpreadsheetProcessor._extract_csv(file_path)
        return SpreadsheetProcessor._extract_xlsx(file_path)

    @staticmethod
    def _extract_csv(file_path: str) -> Dict[str, Any]:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            grid = [row for row in reader if any(cell.strip() for cell in row)]

        if not grid:
            return {"markdown": "*(Empty CSV)*", "sheets": [], "tables": []}

        header = grid[0]
        rows = grid[1:]
        header_line = "| " + " | ".join(header) + " |"
        sep_line = "| " + " | ".join(["---"] * max(1, len(header))) + " |"
        row_lines = ["| " + " | ".join(r) + " |" for r in rows]
        md_tab = "\n".join([header_line, sep_line] + row_lines)

        table_info = {
            "id": "csv_table_1",
            "sheet_name": "Sheet1",
            "markdown": md_tab,
            "num_rows": len(grid),
            "num_cols": len(header),
        }

        return {
            "markdown": f"## CSV Data\n\n{md_tab}",
            "sheets": [{"name": "Sheet1", "rows": len(grid), "cols": len(header)}],
            "tables": [table_info],
        }

    @staticmethod
    def _extract_xlsx(file_path: str) -> Dict[str, Any]:
        import openpyxl

        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        all_markdown: List[str] = []
        tables_data: List[Dict[str, Any]] = []
        sheets_meta: List[Dict[str, Any]] = []

        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                grid: List[List[str]] = []

                for row in ws.iter_rows(values_only=True):
                    row_cells = [str(c or "").replace("\n", " ").strip() for c in row]
                    if any(row_cells):
                        grid.append(row_cells)

                if not grid:
                    continue

                header = grid[0]
                rows = grid[1:]
                header_line = "| " + " | ".join(header) + " |"
                sep_line = "| " + " | ".join(["---"] * max(1, len(header))) + " |"
                row_lines = ["| " + " | ".join(r) + " |" for r in rows]
                md_tab = "\n".join([header_line, sep_line] + row_lines)

                tables_data.append({
                    "id": f"sheet_{sheet_name}",
                    "sheet_name": sheet_name,
                    "markdown": md_tab,
                    "num_rows": len(grid),
                    "num_cols": len(header),
                })
                sheets_meta.append({"name": sheet_name, "rows": len(grid), "cols": len(header)})
                all_markdown.append(f"## Sheet: {sheet_name}\n\n{md_tab}")

        finally:
            wb.close()

        full_md = "\n\n---\n\n".join(all_markdown) if all_markdown else "*(Empty Spreadsheet)*"
        return {
            "markdown": full_md,
            "sheets": sheets_meta,
            "tables": tables_data,
        }


class PresentationProcessor:
    """Extracts slide text, titles, shape tables, and embedded images from PowerPoint (.pptx) files."""

    @staticmethod
    def extract_presentation(file_path: str) -> Dict[str, Any]:
        """Parses a .pptx presentation slide-by-slide."""
        import pptx

        prs = pptx.Presentation(file_path)
        slides_data: List[Dict[str, Any]] = []
        all_markdown: List[str] = []
        embedded_images: List[bytes] = []

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1
            slide_texts: List[str] = []
            title_text = ""
            slide_images: List[bytes] = []

            def _collect_shape_images(shapes_col) -> None:
                for s in shapes_col:
                    if hasattr(s, "image"):
                        try:
                            blob = s.image.blob
                            if blob:
                                slide_images.append(blob)
                                embedded_images.append(blob)
                        except Exception:
                            pass
                    if hasattr(s, "shapes"):
                        _collect_shape_images(s.shapes)

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        txt = paragraph.text.strip()
                        if txt:
                            if not title_text and (shape == slide.shapes.title or "title" in getattr(shape, "name", "").lower()):
                                title_text = txt
                            else:
                                slide_texts.append(txt)

                if shape.has_table:
                    grid: List[List[str]] = []
                    for row in shape.table.rows:
                        grid.append([cell.text.replace("\n", " ").strip() for cell in row.cells])
                    if grid:
                        header = grid[0]
                        rows = grid[1:]
                        header_line = "| " + " | ".join(header) + " |"
                        sep_line = "| " + " | ".join(["---"] * max(1, len(header))) + " |"
                        row_lines = ["| " + " | ".join(r) + " |" for r in rows]
                        slide_texts.append("\n".join([header_line, sep_line] + row_lines))

            # Collect images from shapes and groups
            _collect_shape_images(slide.shapes)

            # Fallback to slide relationships if no images found via shapes
            if not slide_images:
                try:
                    for rel in slide.part.rels.values():
                        if "image" in getattr(rel, "target_ref", "").lower():
                            blob = rel.target_part.blob
                            if blob:
                                slide_images.append(blob)
                                embedded_images.append(blob)
                except Exception:
                    pass

            header_md = f"## Slide {slide_num}: {title_text}" if title_text else f"## Slide {slide_num}"
            body_md = "\n\n".join(slide_texts)
            slide_md = f"{header_md}\n\n{body_md}" if body_md else header_md

            slides_data.append({
                "slide_number": slide_num,
                "title": title_text,
                "text": body_md,
                "markdown": slide_md,
                "embedded_images": slide_images,
            })
            all_markdown.append(slide_md)

        return {
            "markdown": "\n\n---\n\n".join(all_markdown),
            "slides": slides_data,
            "embedded_images": embedded_images,
            "total_slides": len(prs.slides),
        }
